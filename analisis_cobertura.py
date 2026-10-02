#!/usr/bin/env python3

import pandas as pd
import numpy as np
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Noemí Toro Barrios | TFM - Revisión de herramientas y recursos para la anotación y priorización de variantes en regiones no codificantes

#Script para realizar el Perfil de capacidad de anotación de las herramientas.

#Tablas que quiero analizar

input_files = [
    "megatabla_regulación_completa_cadd_gerp.tsv",
    "megatabla_resultados_final_splicing.tsv"
]

id_col = "Variante"
consequence_col = "Consecuencia"
clinvar_col = "Clasificación Clinvar"

NA_values = [
    "No anotada",
    "No anotado",
    "No concluyente"
]  #Modificar todos los scripts si tengo tiempo para unificar la salida


#defino función para leer las tablas. Lo hago de forma que pueda cargar más de una tabla para posibles usos futuros del script
columnas_fijas = [id_col, consequence_col, clinvar_col]


def cargar_tablas(rutas, sep, id_col, como="inner"):  #Como sirve para decir que si todo o lo común

    lista_tablas = [
        pd.read_csv(ruta, sep="\t")
        for ruta in rutas
    ]

    if len(lista_tablas) == 1:
        return lista_tablas[0]

    tabla_fusionada = lista_tablas[0]

    for indice, tabla_extra in enumerate(
        lista_tablas[1:],
        start=2
    ):

        columnas_compartidas = []

        for columna in tabla_extra.columns:
            if columna in tabla_fusionada.columns:
                if columna != id_col:
                    columnas_compartidas.append(columna)

        tabla_fusionada = tabla_fusionada.merge(
            tabla_extra,
            on=id_col,
            how=como,
            suffixes=("", f"_tabla{indice}")
        )

        for columna in columnas_compartidas:

            columna_duplicada = f"{columna}_tabla{indice}"

            if columna_duplicada not in tabla_fusionada.columns:
                continue

            tabla_fusionada[columna] = tabla_fusionada[columna].fillna(
                tabla_fusionada[columna_duplicada]
            )

            tabla_fusionada = tabla_fusionada.drop(
                columns=[columna_duplicada]
            )

    return tabla_fusionada


tabla_principal = cargar_tablas(
    input_files,
    "\t",
    id_col,
    como="inner"
)

columna_herramientas = (
    tabla_principal
    .drop(columns=columnas_fijas)
    .columns
    .tolist()
)


def obtener_nombre(nombre_columna):

    return re.sub(
        r"\s*\([^)]*\)\s*$",
        "",
        nombre_columna
    ).strip()


tabla_principal[columna_herramientas] = (
    tabla_principal[columna_herramientas]
    .replace(NA_values, np.nan)
)


def catalogar_mutacion(variante_texto):

    cromosoma, posicion, referencia, alternativa = (
        str(variante_texto).split(":")
    )

    #Categorizo las variantes en INDELS o SNVs dependiendo de si la referencia y la alternativa es igual a 1.

    if len(referencia) == 1 and len(alternativa) == 1:
        return "SNV"

    return "INDEL"


tabla_principal["Tipo_variante"] = (
    tabla_principal[id_col].apply(catalogar_mutacion)
)


grupo_herramientas = {}

for columna in columna_herramientas:

    nombre_base = obtener_nombre(columna)

    grupo_herramientas.setdefault(
        nombre_base,
        []
    ).append(columna)


def calcular_resumen(datos):

    fila_resultado = {
        "N_variantes": len(datos)
    }

    total_grupo = len(datos)

    for herramienta, columna_agrupadas in grupo_herramientas.items():

        anotada = (
            datos[columna_agrupadas]
            .notna()
            .any(axis=1)
        )

        cantidad_anotada = anotada.sum()

        porcentaje = (
            cantidad_anotada / len(datos) * 100
            if len(datos) > 0
            else 0.0
        )

        fila_resultado[herramienta] = (
            f"{porcentaje:.1f}% ({cantidad_anotada})"
        )

    return fila_resultado


filas_generadas = []
etiquetas_filas = []


#Creo una copia de la tabla para separar las consecuencias múltiples
tabla_por_consecuencia = tabla_principal.copy()

tabla_por_consecuencia[consequence_col] = (
    tabla_por_consecuencia[consequence_col]
    .fillna("Sin_Dato")
    .astype(str)
    .str.split("|")
)

tabla_por_consecuencia = (
    tabla_por_consecuencia
    .explode(consequence_col)
)

tabla_por_consecuencia[consequence_col] = (
    tabla_por_consecuencia[consequence_col]
    .str.strip()
)


#Separo por consecuencia molecular y luego por SNV/INDELS

for consecuencia, datos_consecuencia in (
    tabla_por_consecuencia.groupby(consequence_col)
):

    for tipo, datos_tipo in datos_consecuencia.groupby(
        "Tipo_variante"
    ):

        filas_generadas.append(
            calcular_resumen(datos_tipo)
        )

        etiquetas_filas.append(
            (consecuencia, tipo)
        )

    #Subtotal consecuencia

    filas_generadas.append(
        calcular_resumen(datos_consecuencia)
    )

    etiquetas_filas.append(
        (consecuencia, "TOTAL")
    )


for tipo, datos_tipo_general in tabla_principal.groupby(
    "Tipo_variante"
):

    filas_generadas.append(
        calcular_resumen(datos_tipo_general)
    )

    etiquetas_filas.append(
        ("TOTAL (todas consecuencias)", tipo)
    )


filas_generadas.append(
    calcular_resumen(tabla_principal)
)

etiquetas_filas.append(
    ("TOTAL GENERAL", "TOTAL")
)


carpeta_salida = "perfil_deteccion"

os.makedirs(
    carpeta_salida,
    exist_ok=True
)


tabla_final = pd.DataFrame(
    filas_generadas,
    index=pd.MultiIndex.from_tuples(
        etiquetas_filas,
        names=[
            consequence_col,
            "Tipo_variante"
        ]
    ),
)


ruta_excel = (
    f"{carpeta_salida}/perfil_deteccion_herramientas.xlsx"
)

tabla_final.to_excel(ruta_excel)


wb = openpyxl.load_workbook(ruta_excel)
ws = wb.active


# Definimos estilos limpios y profesionales
fuente_cabecera = Font(
    name="Calibri",
    size=11,
    bold=True,
    color="FFFFFF"
)

relleno_cabecera = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid"
)

borde_fino = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3")
)

alineacion_centrada = Alignment(
    horizontal="center",
    vertical="center"
)

alineacion_izquierda = Alignment(
    horizontal="left",
    vertical="center"
)


# Aplicar formato a las cabeceras
for row in ws.iter_rows(
    min_row=1,
    max_row=ws.max_row
):

    for cell in row:

        cell.border = borde_fino

        if cell.row <= 2:

            cell.font = fuente_cabecera
            cell.fill = relleno_cabecera
            cell.alignment = alineacion_centrada

        else:

            if cell.column == 1:
                cell.alignment = alineacion_izquierda
            else:
                cell.alignment = alineacion_centrada


# Ajustar automáticamente el ancho de las columnas
for col in ws.columns:

    max_length = 0
    col_letter = get_column_letter(
        col[0].column
    )

    for cell in col:

        try:

            if cell.value:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        except:
            pass

    ws.column_dimensions[col_letter].width = max(
        max_length + 4,
        12
    )


# Guardamos los cambios de formato
wb.save(ruta_excel)


print(
    f"\nAnálisis de perfil de detección terminado: "
    f"{ruta_excel}"
)


tabla_principal.to_csv(
    "tabla_fusionada.tsv",
    sep="\t",
    index=False
)